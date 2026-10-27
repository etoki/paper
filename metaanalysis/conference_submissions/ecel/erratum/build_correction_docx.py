"""Build correction_EL-095.docx from correction_EL-095.md.

Handles the subset of Markdown the correction uses: '#'/'##' headings,
paragraphs, '- ' bullets, pipe tables, **bold** and *italic*.

Usage:
    python metaanalysis/conference_submissions/ecel/erratum/build_correction_docx.py [SRC.md OUT.docx]

With no arguments it builds the ECEL correction; other papers pass their
own Markdown source and output path.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
SRC = HERE / "correction_EL-095.md"
OUT = HERE / "correction_EL-095.docx"

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def add_inline(par, text: str, size: int = 10) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*"):
            run = par.add_run(part[1:-1])
            run.italic = True
        else:
            run = par.add_run(part)
        run.font.size = Pt(size)


def add_table(doc, lines: list[str]) -> None:
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")]
            for ln in lines if not re.fullmatch(r"\|[-| ]+\|", ln.strip())]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, cells in enumerate(rows):
        for j, text in enumerate(cells):
            par = table.cell(i, j).paragraphs[0]
            add_inline(par, text, size=9)
            if i == 0:
                for run in par.runs:
                    run.bold = True


def main() -> None:
    src, out = (Path(a) for a in sys.argv[1:3]) if len(sys.argv) == 3 else (SRC, OUT)
    doc = Document()
    lines = src.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            doc.add_paragraph()
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("- "):
            add_inline(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            add_inline(doc.add_paragraph(), line)
        i += 1
    doc.save(str(out))
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
